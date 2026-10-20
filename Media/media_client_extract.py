#!/usr/bin/env python3
"""
Client-side media extract for AssetFuture Developer API (Documents V2).

Requires only an APIM subscription key. Configure the API URL in media_extract.settings.json:
  1. POST /v2/Documents/metadata (paginated)
  2. GET  /v2/Documents/{mediaId}/file for each photo
  3. Builds index.xlsx + unmapped-areas.csv
  4. Creates the deliverable ZIP (outputPath + outputFileName in media_extract.settings.json)

Usage:
  python media_client_extract.py <subscription-key>
  python media_client_extract.py   # when Ocp-Apim-Subscription-Key env var is set

Settings (media_extract.settings.json):
  baseUrl          - Media API root URL
  filter           - Optional metadata query (space, areaIdPrefix) — see FILTER CONFIG below
  outputPath       - Folder for photos, index.xlsx, metadata.json (supports {timestamp})
  outputFileName   - Optional. ZIP name inside outputPath (default: media-extract-{timestamp}.zip)
"""

# ---------------------------------------------------------------------------
# FILTER CONFIG (media_extract.settings.json)
# ---------------------------------------------------------------------------
# Sent as the POST /v2/Documents/metadata JSON body. Use {} for all project photos.
#
# All photos:
#   "filter": {}
#
# Area photos only:
#   "filter": { "space": "Area" }
#
# Item photos in areas whose number starts with "1.2":
#   "filter": { "space": "Item", "areaIdPrefix": "1.2" }
#
# Any space, limited by area number prefix:
#   "filter": { "areaIdPrefix": "Bldg-A" }
#
# space — optional; one of: Area, Item, Issue
# areaIdPrefix — optional; area number prefix (letters, digits, . _ -)
#
# See media_extract.settings.json → _comments.filter for examples in config.
# ---------------------------------------------------------------------------

from __future__ import annotations

import csv
import json
import os
import sys
import time
import zipfile
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from openpyxl.styles import Font

SCRIPT_DIR = Path(__file__).resolve().parent
SETTINGS_FILE = SCRIPT_DIR / "media_extract.settings.json"

INDEX_HEADERS = [
    "Type",
    "AreaId",
    "AreaPath",
    "ItemId",
    "IssueId",
    "Title",
    "FileName",
    "Link",
]

UNMAPPED_HEADERS = [
    "MediaId",
    "ExternalAreaId",
    "ExternalAreaMongoId",
    "ExternalAreaSqlId",
]


def log(message: str) -> None:
    print(message, flush=True)


def load_settings() -> dict[str, Any]:
    if not SETTINGS_FILE.exists():
        return {
            "baseUrl": "http://localhost:5005",
            "filter": {},
            "outputPath": "output/media-extract-{timestamp}",
            "pageSize": 1000,
            "progressEvery": 10,
        }

    with SETTINGS_FILE.open(encoding="utf-8") as handle:
        return json.load(handle)


def expand_timestamp(value: str, timestamp: str) -> str:
    return value.replace("{timestamp}", timestamp)


DEFAULT_OUTPUT_FILE_NAME = "media-extract-{timestamp}.zip"


def resolve_output_paths(settings: dict[str, Any]) -> tuple[Path, Path]:
    timestamp = time.strftime("%Y%m%d-%H%M%S")
    output_path = expand_timestamp(
        settings.get("outputPath") or "output/media-extract-{timestamp}",
        timestamp,
    )
    output_file_name = (settings.get("outputFileName") or "").strip()
    if not output_file_name:
        output_file_name = DEFAULT_OUTPUT_FILE_NAME
    output_file_name = expand_timestamp(output_file_name, timestamp)

    output_dir = Path(output_path)
    if not output_dir.is_absolute():
        output_dir = (SCRIPT_DIR / output_dir).resolve()

    zip_name = Path(output_file_name).name
    if not zip_name.lower().endswith(".zip"):
        zip_name = f"{zip_name}.zip"

    return output_dir, output_dir / zip_name


def is_api_reachable(base_url: str, timeout_seconds: float = 8.0) -> bool:
    try:
        response = requests.get(
            f"{base_url.rstrip('/')}/swagger/v1/swagger.json",
            timeout=timeout_seconds,
        )
        if not response.ok:
            return False
        title = response.json().get("info", {}).get("title", "")
        return "AssetFuture.API.Media" in title
    except requests.RequestException:
        return False


def resolve_base_url(settings: dict[str, Any]) -> str:
    base_url = (settings.get("baseUrl") or "").strip().rstrip("/")
    if not base_url:
        raise SystemExit(f"baseUrl is required in {SETTINGS_FILE.name}.")

    if not is_api_reachable(base_url):
        raise SystemExit(
            f"Media API is not reachable at {base_url}. "
            f"Check baseUrl in {SETTINGS_FILE.name} and ensure the service is running."
        )

    return base_url


def resolve_metadata_filter(settings: dict[str, Any]) -> dict[str, Any]:
    raw = settings.get("filter")
    if raw is None:
        return {}
    if not isinstance(raw, dict):
        raise SystemExit(f"filter must be a JSON object in {SETTINGS_FILE.name}.")

    body: dict[str, Any] = {}

    space = raw.get("space")
    if space is not None:
        space = str(space).strip()
        if space:
            body["space"] = space

    area_id_prefix = raw.get("areaIdPrefix")
    if area_id_prefix is not None:
        area_id_prefix = str(area_id_prefix).strip()
        if area_id_prefix:
            body["areaIdPrefix"] = area_id_prefix

    return body


def resolve_subscription_key(argv: list[str]) -> str:
    if len(argv) > 1 and argv[1].strip():
        return argv[1].strip()

    key = os.environ.get("Ocp-Apim-Subscription-Key", "").strip()
    if key:
        return key

    print(
        "Error: subscription key required.\n"
        "  python media_client_extract.py <subscription-key>\n"
        "  or set Ocp-Apim-Subscription-Key environment variable.",
        file=sys.stderr,
    )
    sys.exit(1)


def api_headers(subscription_key: str) -> dict[str, str]:
    return {
        "Ocp-Apim-Subscription-Key": subscription_key,
        "Content-Type": "application/json",
        "Accept": "application/json",
    }


def fetch_all_metadata(
    session: requests.Session,
    base_url: str,
    headers: dict[str, str],
    page_size: int,
    metadata_filter: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int]:
    url = f"{base_url}/v2/Documents/metadata"
    body = metadata_filter

    items: list[dict[str, Any]] = []
    unmapped: list[dict[str, Any]] = []
    offset = 0
    total_matched = 0
    page_number = 0

    while True:
        page_number += 1
        params = {"offset": offset, "limit": page_size}
        log(f"  metadata page {page_number}: offset={offset}, limit={page_size}...")
        started = time.time()
        response = session.post(url, headers=headers, json=body, params=params, timeout=300)
        response.raise_for_status()
        page = response.json()
        elapsed = time.time() - started

        total_matched = int(page.get("totalMatched", 0))
        page_items = page.get("items", [])
        page_unmapped = page.get("unmapped", [])
        items.extend(page_items)
        unmapped.extend(page_unmapped)
        log(
            f"  metadata page {page_number}: +{len(page_items)} mapped, "
            f"+{len(page_unmapped)} unmapped ({elapsed:.1f}s) — total {len(items)}/{total_matched}"
        )

        if not page.get("hasMore", False):
            break

        offset += int(page.get("limit", page_size))

    return items, unmapped, total_matched


def download_photo(
    session: requests.Session,
    base_url: str,
    subscription_key: str,
    media_id: str,
    destination: Path,
) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() and destination.stat().st_size > 0:
        return

    url = f"{base_url}/v2/Documents/{media_id}/file"
    headers = {"Ocp-Apim-Subscription-Key": subscription_key, "Accept": "image/jpeg"}

    with session.get(url, headers=headers, stream=True, timeout=600) as response:
        response.raise_for_status()
        with destination.open("wb") as handle:
            for chunk in response.iter_content(chunk_size=1024 * 1024):
                if chunk:
                    handle.write(chunk)


def write_unmapped_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return

    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=UNMAPPED_HEADERS)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "MediaId": row.get("mediaId", ""),
                    "ExternalAreaId": row.get("externalAreaId", "") or "",
                    "ExternalAreaMongoId": row.get("externalAreaMongoId", "") or "",
                    "ExternalAreaSqlId": row.get("externalAreaSqlId", "") or "",
                }
            )


def build_index_xlsx(path: Path, items: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Media"

    for column, header in enumerate(INDEX_HEADERS, start=1):
        sheet.cell(row=1, column=column, value=header)

    hyperlink_font = Font(color="0563C1", underline="single")

    for row_index, item in enumerate(items, start=2):
        relative_path = item.get("relativePath", "")
        values = [
            item.get("type", ""),
            item.get("areaId", ""),
            item.get("areaPath", ""),
            item.get("itemId", ""),
            item.get("issueId", ""),
            item.get("title", ""),
            item.get("fileName", ""),
            relative_path,
        ]

        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value if column < 8 else "See Photo")
            if column == 8 and relative_path:
                cell.hyperlink = relative_path
                cell.font = hyperlink_font

    workbook.save(path)


def create_zip(source_dir: Path, zip_path: Path) -> None:
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_STORED) as archive:
        photos_dir = source_dir / "photos"
        if photos_dir.exists():
            for file_path in sorted(photos_dir.rglob("*")):
                if file_path.is_file():
                    archive.write(file_path, file_path.relative_to(source_dir).as_posix())

        for name in ("index.xlsx", "unmapped-areas.csv"):
            file_path = source_dir / name
            if file_path.exists():
                archive.write(file_path, name)


def main() -> None:
    settings = load_settings()
    subscription_key = resolve_subscription_key(sys.argv)
    base_url = resolve_base_url(settings)
    page_size = int(settings.get("pageSize", 1000))
    progress_every = max(int(settings.get("progressEvery", 10)), 1)
    metadata_filter = resolve_metadata_filter(settings)

    output_dir, zip_path = resolve_output_paths(settings)
    output_dir.mkdir(parents=True, exist_ok=True)

    headers = api_headers(subscription_key)
    session = requests.Session()

    log(f"Starting client extract -> {output_dir}")
    log(f"API: {base_url}")
    if metadata_filter:
        log(f"Filter: {json.dumps(metadata_filter)}")
    else:
        log("Filter: (none — all project photos)")
    log("Fetching metadata...")
    items, unmapped, total_matched = fetch_all_metadata(
        session, base_url, headers, page_size, metadata_filter
    )
    log(f"Matched {total_matched} photo(s); {len(items)} mapped, {len(unmapped)} unmapped.")

    metadata_path = output_dir / "metadata.json"
    metadata_path.write_text(
        json.dumps(
            {
                "totalMatched": total_matched,
                "mappedCount": len(items),
                "unmappedCount": len(unmapped),
                "items": items,
                "unmapped": unmapped,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    log(f"Saved {metadata_path}")

    log(f"Downloading {len(items)} photo(s)...")
    for index, item in enumerate(items, start=1):
        media_id = item["mediaId"]
        relative_path = item["relativePath"]
        destination = output_dir / Path(relative_path.replace("/", os.sep))
        if index == 1 or index % progress_every == 0 or index == len(items):
            log(f"  download {index}/{len(items)}: {media_id}")
        download_photo(session, base_url, subscription_key, media_id, destination)

    log("Building index.xlsx...")
    build_index_xlsx(output_dir / "index.xlsx", items)

    unmapped_path = output_dir / "unmapped-areas.csv"
    write_unmapped_csv(unmapped_path, unmapped)
    if unmapped:
        log(f"Wrote {unmapped_path} ({len(unmapped)} row(s))")

    log(f"Creating ZIP {zip_path}...")
    create_zip(output_dir, zip_path)
    log(f"Done. ZIP: {zip_path}")


if __name__ == "__main__":
    main()
